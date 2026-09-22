from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = r"D:\shoe(1)\shoe\晶格过渡相关文献调研.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=110, start=110, bottom=110, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_text(cell, text, bold=False, size=8.6):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return p


direct = [
    ("Gao, D. et al. (2024)", "Topology-aware blending method for implicit heterogeneous porous model design. Computer-Aided Design.", "10.1016/j.cad.2024.103782", "不同多孔隐式场的融合。以持久同调刻画并优化融合后错误孔洞、孤立分量和断连。", "最接近现有 Ramp 融合与拓扑约束修正。可将启发式局部修正升级为可验证的连通性和孔洞约束。"),
    ("Tian et al. (2024)", "Continuous transitions of triply periodic minimal surfaces. Additive Manufacturing.", "10.1016/j.addma.2024.104105", "研究不同 TPMS 形态之间的连续过渡及其结构表现。", "直接对应 G、D 与其他 TPMS 单元的连续过渡，是当前系统核心功能的主要对照文献。"),
    ("Li et al. (2025)", "Parametric Design and Mechanical Behavior of Hybrid TPMS Lattice Structures Based on Sigmoid Function. Advanced Engineering Materials.", "10.1002/adem.202402360", "以 sigmoid 函数定义混合 TPMS 的过渡规律，并分析力学行为。", "与现有 sigmoid Ramp 的宽度、锐度和场值混合逻辑最接近，可用于参数与力学性能的验证。"),
    ("Moitra et al. (2026)", "Strengthening the transition zone of an additively manufactured hybrid TPMS structure for bone scaffold applications. International Journal of Solids and Structures.", "10.1016/j.ijsolstr.2026.114025", "围绕混合 TPMS 的过渡带强化开展设计和力学评估。", "说明过渡区目标不应止于几何连续，还应评估局部强度、可制造性和性能突变。"),
    ("Mathiazhagan et al. (2026)", "Modal analysis of 3D-printed PA12 TPMS lattices incorporating sigmoid transition zones. Composites Part C.", "10.1016/j.jcomc.2026.100765", "将 sigmoid 过渡区纳入 3D 打印 TPMS 晶格的模态分析。", "可作为未来把过渡函数、过渡宽度与动力学性能关联的工程验证参考。"),
    ("Letov and Zhao (2022)", "Beam-Based Lattice Topology Transition With Function Representation. Journal of Mechanical Design.", "10.1115/1.4055950", "以函数表示 F-Rep 描述梁式晶格，并连续控制参数以实现不同梁拓扑间转变。", "虽不是 TPMS，但与任意自定义晶胞的函数表达及场级连续过渡高度相关。"),
    ("Yoo (2013)", "Heterogeneous porous scaffold design using the continuous transformations of triply periodic minimal surface models. International Journal of Precision Engineering and Manufacturing.", "10.1007/s12541-013-0234-4", "早期的 TPMS 连续变换异质多孔支架设计。", "可作为由两端晶胞构成连续中间形态的经典技术背景。"),
    ("Yang et al. (2026)", "Pyramid connection method for constructing graded triply periodic minimal surface lattices. Thin-Walled Structures.", "10.1016/j.tws.2025.114253", "采用显式连接构造处理梯度 TPMS 单元间的衔接。", "与系统的连续场混合形成对照：该文偏几何连接，当前系统偏隐式场融合。"),
]

field = [
    ("Hong, Elber and Kim (2023)", "Implicit Functionally Graded Conforming Microstructures. Computer-Aided Design.", "10.1016/j.cad.2023.103548", "以三变量映射 T(u,v,w) 将物理空间反求至参数 UVW 域；填充隐式 tile，并用连续标量场保证相邻单元连续。", "与 Cell Map、UVW、权威 SDF 和采样重建架构高度一致；可支撑自由形或非规则 Cell Map 的演进。"),
    ("Hu and Lin (2021)", "Heterogeneous porous scaffold generation using trivariate B-spline solids and triply periodic minimal surfaces. Graphical Models.", "10.1016/j.gmod.2021.101105", "基于三变量 B-spline 实体构造异质 TPMS 多孔支架。", "为曲面贴合、空间变尺度、变方向的参数化 Cell Map 提供方法参照。"),
    ("Yoo (2012)", "Heterogeneous minimal surface porous scaffold design using the distance field and radial basis functions. Medical Engineering and Physics.", "10.1016/j.medengphy.2012.03.009", "将带属性的控制点经薄板 RBF 插值得到连续标量场，再控制 TPMS 局部孔隙率。", "最适合场驱动设计：应力、温度或用户控制点场可驱动壁厚、单元尺度、过渡权重和晶胞类型。"),
    ("Yoo (2011)", "Porous scaffold design using the distance field and triply periodic minimal surface models. Biomaterials.", "10.1016/j.biomaterials.2011.07.019", "把设计对象距离场与 TPMS 隐式表达组合，以适配任意外形及多孔填充。", "是设计域 SDF 与晶格场求交、裁剪这一权威隐式体路线的重要基础。"),
]

ai = [
    ("Zheng et al. (2023)", "Unifying the design space and optimizing linear and nonlinear truss metamaterials by generative modeling. Nature Communications.", "10.1038/s41467-023-42068-x", "以图表示晶胞，用 VAE 构建连续潜在空间，并加入性质预测与目标优化。", "未来可用当前 Ramp 权重插值潜变量 z(t)=(1-w)zA+wzB，再解码出更自然的中间晶胞。"),
    ("Kim et al. (2024)", "Simple arithmetic operation in latent space can generate a novel three-dimensional graph metamaterials. npj Computational Materials.", "10.1038/s41524-024-01430-3", "以解耦潜在表示和生成模型，通过潜在空间算术生成新图晶格及梯度材料。", "对应从场值插值升级为学习到的晶胞生成参数插值，是 AI 过渡路线的代表。"),
    ("Abu-Mualla and Huang (2025)", "Inverse design of curved mechanical metamaterials with geometric AI. npj Metamaterials.", "10.1038/s44455-025-00005-6", "几何 AI 驱动的曲面机械超材料反设计。", "适合后续结合曲面 Cell Map、目标性能场与自适应晶胞生成。"),
    ("Viswanath et al. (2024)", "Designing a TPMS metamaterial via deep learning and topology optimization. Frontiers in Mechanical Engineering.", "10.3389/fmech.2024.1417606", "结合深度学习与拓扑优化进行 TPMS 超材料设计。", "可作为场驱动 TPMS 性能代理模型与优化闭环的工程参考。"),
    ("Yeo et al. (2024)", "Hybrid TPMS-based architectured materials for enhanced specific stiffness using data-driven design. Materials and Design.", "10.1016/j.matdes.2024.113523", "数据驱动地筛选混合 TPMS 构型和性能。", "可用于建立晶胞组合、过渡参数和力学性能之间的数据集及推荐机制。"),
]

validation = [
    ("Zhang et al. (2024)", "A study of energy absorption properties of Heteromorphic TPMS and Multi-morphology TPMS under quasi-static compression. Thin-Walled Structures.", "10.1016/j.tws.2024.112519", "比较异形和多形态 TPMS 的准静态压缩及吸能性能。", "用于评估过渡方案在吸能和平台应力层面的工程收益。"),
    ("Xi et al. (2023)", "Multi-morphology TPMS structures with multi-stage yield stress platform and multi-level energy absorption. Engineering Structures.", "10.1016/j.engstruct.2023.116733", "研究多形态 TPMS 的多阶段屈服平台和多级吸能。", "用于定义混合晶格的性能评价指标。"),
    ("Yin et al. (2025)", "Enhanced energy absorption characteristics of TPMS lattice structures with linear and circular hybrid designs. Engineering Structures.", "10.1016/j.engstruct.2025.120759", "比较线性和圆形混合设计下 TPMS 晶格的吸能特性。", "与系统的平面和任意场对象过渡区域定义直接相关。"),
]

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.1); sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.0); sec.right_margin = Cm(2.0)
styles = doc.styles
styles['Normal'].font.name = 'Microsoft YaHei'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
styles['Normal'].font.size = Pt(10)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.35
for style_name, size in [('Title', 22), ('Heading 1', 15), ('Heading 2', 12)]:
    s = styles[style_name]; s.font.name = 'Microsoft YaHei'; s._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei'); s.font.size = Pt(size); s.font.color.rgb = RGBColor(0,0,0)

title = doc.add_paragraph(style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('晶格过渡相关文献调研').font.color.rgb = RGBColor(0,0,0)
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('面向隐式场融合 Cell Map 参数化与场驱动晶格设计'); r.font.size = Pt(11); r.font.color.rgb = RGBColor(89,89,89)

doc.add_heading('调研结论', level=1)
doc.add_paragraph('当前系统的技术路线可归纳为：在共享 Cell Map 和 UVW 参数空间中，以权威 SDF 表达晶胞和设计域；由平面或任意场对象驱动 Ramp 权重，对两类晶胞场进行连续融合；再经渲染或 STL 重建采样输出。文献中最直接可借鉴的方向是拓扑感知隐式融合，其次是三变量参数化微结构、距离场与 RBF 驱动梯度设计，以及潜在空间中的晶胞生成与插值。')
p = doc.add_paragraph(); p.add_run('建议精读顺序：').bold = True; p.add_run('Gao 2024（拓扑正确性） -> Tian 2024（TPMS 过渡） -> Hong 2023（UVW 参数化） -> Letov and Zhao 2022（函数表示） -> Zheng 2023（潜在空间）。')

doc.add_heading('检索与核验说明', level=1)
doc.add_paragraph('本报告通过 paper-search-mcp 检索 CrossRef、OpenAlex 和 OpenAIRE，并以 DOI 与出版社元数据核验题名、作者、期刊和发表年份。2026 年条目以在线出版记录为准，引用积累可能仍有限；其真实性不等同于影响力已经充分建立。')

def add_section(title, rows):
    doc.add_heading(title, level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    table.autofit = False
    widths = [Cm(4.0), Cm(5.4), Cm(4.3), Cm(4.0)]
    heads = ['文献', '工作与方法', '与当前项目的关系', '核验']
    for i, cell in enumerate(table.rows[0].cells):
        cell.width = widths[i]; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_shading(cell, '1F4E78'); set_cell_margins(cell)
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run = p.add_run(heads[i]); run.bold = True; run.font.size = Pt(9); run.font.color.rgb = RGBColor(255,255,255)
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement('w:tblHeader')
    repeat_header.set(qn('w:val'), 'true')
    header_pr.append(repeat_header)
    for n, (author, paper, doi, method, relevance) in enumerate(rows):
        cells = table.add_row().cells
        row_pr = table.rows[-1]._tr.get_or_add_trPr()
        cant_split = OxmlElement('w:cantSplit')
        row_pr.append(cant_split)
        for i, cell in enumerate(cells):
            cell.width = widths[i]; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cell)
            if n % 2: set_cell_shading(cell, 'F2F6FA')
        p = add_text(cells[0], author + '\n', True)
        rr = p.add_run(paper); rr.font.size = Pt(8.3); rr.italic = True
        add_text(cells[1], method)
        add_text(cells[2], relevance)
        p = cells[3].paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.15
        rr = p.add_run('CrossRef / 出版社 DOI\n'); rr.font.size = Pt(8.3); rr.bold = True
        hyperlink(p, doi, 'https://doi.org/' + doi)
    doc.add_paragraph()

add_section('一 直接面向晶格过渡与拓扑正确性的研究', direct)
add_section('二 场驱动与参数化隐式微结构研究', field)
add_section('三 潜在空间与数据驱动晶格生成研究', ai)
add_section('四 混合 TPMS 工程性能验证研究', validation)

doc.add_heading('对当前系统的下一步建议', level=1)
for text in [
    '将过渡区连通性由经验式后处理逐步升级为可计算的拓扑目标，优先检测孤立实体分量、断连主通路和非预期孔洞。',
    '保持场对象定义过渡区域的架构，并扩展 RBF 或仿真场驱动接口，使场可同时控制晶胞类型、尺寸、壁厚和融合权重。',
    '以共享 UVW 与 Cell Map 为核心，探索曲线和自由形参数化映射，避免复杂设计域中仅靠规则周期复制导致的边界失配。',
    '将潜在空间晶胞生成作为中长期路线：先建立可制造晶胞数据集与性能标签，再将 Ramp 作为潜变量空间路径的控制器。',
]:
    doc.add_paragraph(text, style='List Bullet')

footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run('晶格过渡相关文献调研'); fr.font.name = 'Microsoft YaHei'; fr.font.size = Pt(8); fr.font.color.rgb = RGBColor(100,100,100)
doc.core_properties.title = '晶格过渡相关文献调研'
doc.core_properties.subject = '隐式场融合与 TPMS 晶格过渡文献整理'
doc.core_properties.author = 'Shoe Lattice Project'
doc.save(OUT)
print(OUT)
